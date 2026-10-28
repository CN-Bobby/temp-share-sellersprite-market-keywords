#!/usr/bin/env python3
"""build_template.py — 生成/重建 assets/keyword_market_report_template.xlsx

生成 13 个 Sheet 的空模板（表头 + 冻结 + 筛选），供 build_excel_report.py --template 使用，
也便于人工查看报告结构。修改模板结构后重新运行本脚本即可。

用法: python build_template.py [--out ../assets/keyword_market_report_template.xlsx]
"""
import argparse
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
FONT_HEADER = Font(color="FFFFFF", bold=True)

SHEETS = {
    "00_市场看板": ["项目", "内容"],
    "01_关键词总表": ["数据月份", "原始关键词", "MCP中文翻译", "校正中文", "标准词族", "需求标签", "流量性质",
                    "流量性质理由", "相关度", "相关度理由", "购买意图", "月搜索量", "月购买量", "购买率",
                    "点击集中度", "转化集中度", "PPC默认", "PPC最低", "PPC最高", "商品数",
                    "1日广告竞品数", "30日广告竞品数", "供需比", "标题密度", "当前趋势资格",
                    "来源标识", "数据质量备注"],
    "02_需求分类汇总": ["需求标签", "关键词数", "代表词", "搜索量中位数", "购买率中位数",
                      "点击集中度中位数", "转化集中度中位数", "当前PPC中位数", "趋势方向", "产品含义"],
    "03_24月搜索量_宽表": ["关键词", "中文翻译", "类型"],
    "04_24月搜索量_长表": ["marketplace", "keyword", "中文翻译", "关键词类型", "month",
                         "searches", "ABA rank", "MoM", "YoY", "seasonality_index"],
    "05_趋势与季节性": ["关键词", "类型", "24个月均值", "24个月中位数", "最新月搜索量", "MoM", "YoY",
                      "近3月增长率", "近6月增长率", "12月同比累计", "变异系数", "峰谷比", "峰值月份",
                      "谷值月份", "连续增长月数", "增长拐点", "两年重复度", "缺月", "生命周期",
                      "进入时间建议"],
    "06_集中度与机会矩阵": ["关键词", "点击集中度", "转化集中度", "象限", "机会类型",
                          "数据证据", "置信度", "解释"],
    "07_产品方向": ["需求板块", "证据关键词", "数据证据", "用户可能问题（推断）", "产品形态", "材料", "功能", "规格",
                   "套装/配件", "产品响应假设", "置信度", "需要验证的数据"],
    "08_ASIN自然广告位矩阵": ["矩阵观测时间", "原始关键词", "中文翻译"],
    "09_ASIN流量占比矩阵": ["矩阵观测时间", "原始关键词", "中文翻译"],
    "10_市场进入与风险": ["项目", "内容"],
    "11_调研范围与调用": ["项目", "值"],
    "12_数据质量": ["类别", "对象", "说明", "是否影响结论"],
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(os.path.dirname(__file__), "..", "assets",
                                                  "keyword_market_report_template.xlsx"))
    args = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)
    for name, header in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = FILL_HEADER
            cell.font = FONT_HEADER
            ws.column_dimensions[get_column_letter(c)].width = 18
        ws.freeze_panes = "A2"
        if name not in ("00_市场看板", "10_市场进入与风险", "11_调研范围与调用"):
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}1"
    out = os.path.abspath(args.out)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    print(f"OK 模板已生成: {out}")


if __name__ == "__main__":
    main()
