import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import LineChart


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def run_script(name, *args):
    env = os.environ.copy()
    env["PYTHONUTF8"] = "1"
    try:
        return subprocess.run(
            [sys.executable, str(SCRIPTS / name), *map(str, args)], check=True,
            capture_output=True, text=True, encoding="utf-8", env=env)
    except subprocess.CalledProcessError as exc:
        raise AssertionError(f"{name} failed\nSTDOUT:\n{exc.stdout}\nSTDERR:\n{exc.stderr}") from exc


def month_range(count=24, start=(2024, 8)):
    out = []
    year, month = start
    for _ in range(count):
        out.append(f"{year:04d}-{month:02d}")
        month += 1
        if month == 13:
            year += 1
            month = 1
    return out


def series_title(series):
    tx = getattr(series, "tx", None)
    return getattr(tx, "v", None) if tx else None


def series_color(series):
    color = series.graphicalProperties.line.solidFill
    return getattr(color, "srgbClr", None)


def series_category_formula(series):
    return series.cat.strRef.f


class V161Tests(unittest.TestCase):
    def test_mode_detection_and_pure_node_matrix(self):
        spec = importlib.util.spec_from_file_location("build_excel_report", SCRIPTS / "build_excel_report.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        self.assertEqual(module.detect_mode({"asins": ["B0AAA"], "node_ids": ["123"]}), "M1")
        self.assertEqual(module.detect_mode({"asins": ["B0AAA", "B0BBB"], "node_ids": []}), "M2")
        self.assertEqual(module.detect_mode({"asins": [], "node_ids": ["123"]}), "M3")
        self.assertEqual(module.detect_mode({"asins": ["B0AAA"], "node_ids": []}), "M4")
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            write_json(wd / "input.json", {
                "groups": [{"marketplace": "US", "asins": [], "node_ids": ["123"]}],
                "options": {"data_month": "202607"},
            })
            run_script("aggregate_asin_keyword_matrix.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "asin_keyword_matrix.json").read_text(encoding="utf-8"))
            self.assertEqual(result["snapshot_type"], "current")
            self.assertEqual(result["scope"], "current_top100")
            self.assertEqual(result["page_limit"], 1)
            self.assertEqual(result["records"], [])

    def test_current_matrix_extracts_page_index_and_ignores_legacy(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            write_json(wd / "input.json", {
                "groups": [{"marketplace": "US", "asins": ["B0AAA", "B0BBB"], "node_ids": []}],
                "options": {"data_month": "202607"},
            })
            write_json(wd / "raw" / "traffic_keyword_B0AAA_p1.json", {
                "_meta": {"asin": "B0AAA", "data_month": "202607", "page": 1},
                "code": "OK", "data": {"items": [{"keyword": "legacy keyword", "trafficPercentage": 0.9}]},
            })
            write_json(wd / "raw" / "traffic_keyword_current_B0AAA_p1.json", {
                "_meta": {"asin": "B0AAA", "marketplace": "US", "observed_at": "2026-08-27T10:00:00+08:00",
                          "page": 1, "page_size": 100, "snapshot_type": "current"},
                "code": "OK", "data": {"total": 99999, "items": [
                    {"keyword": "alpha keyword", "rankPosition": {"page": 1, "index": 36, "position": 36, "updatedTime": 100},
                     "adPosition": {"page": 2, "index": 8, "position": 96, "updatedTime": 100},
                     "trafficPercentage": 0.12, "updatedTime": 100},
                    {"keyword": "alpha keyword", "rankPosition": {"page": 1, "index": 35, "position": 35, "updatedTime": 200},
                     "adPosition": {"page": 2, "index": 7, "position": 95, "updatedTime": 200},
                     "trafficPercentage": 0.13, "updatedTime": 200},
                    {"keyword": "natural only", "rankPosition": {"page": 1, "index": 9, "position": 9},
                     "adPosition": None, "trafficPercentage": 0.05},
                    {"keyword": "ad only", "rankPosition": None,
                     "adPosition": {"page": 1, "index": 4, "position": 4}, "trafficPercentage": None},
                    {"keyword": "blank metrics", "rankPosition": None, "adPosition": None, "trafficPercentage": None},
                ]},
            })
            write_json(wd / "raw" / "traffic_keyword_current_B0BBB_p1.json", {
                "_meta": {"asin": "B0BBB", "page": 1, "page_size": 100},
                "code": "ERROR_UNAUTHORIZED", "message": "quota or authorization rejected",
            })
            run_script("aggregate_asin_keyword_matrix.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "asin_keyword_matrix.json").read_text(encoding="utf-8"))
            by_key = {(r["asin"], r["keyword"]): r for r in result["records"]}
            alpha = by_key[("B0AAA", "alpha keyword")]
            self.assertEqual((alpha["natural_page"], alpha["natural_index"]), (1, 35))
            self.assertEqual((alpha["ad_page"], alpha["ad_index"]), (2, 7))
            self.assertAlmostEqual(alpha["traffic_percentage"], 0.13)
            self.assertNotIn(("B0AAA", "legacy keyword"), by_key)
            self.assertEqual(result["coverage"]["B0AAA"]["pages"], [1])
            self.assertTrue(result["coverage"]["B0AAA"]["request_success"])
            self.assertFalse(result["coverage"]["B0BBB"]["request_success"])
            self.assertIn("授权或额度拒绝", result["coverage"]["B0BBB"]["notes"][0])

    def test_multi_asin_batch_does_not_fake_keyword_coverage(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            write_json(wd / "input.json", {
                "groups": [{"marketplace": "US", "asins": ["B0AAA", "B0BBB"], "node_ids": []}],
                "options": {"data_month": "202607"},
            })
            write_json(wd / "raw" / "traffic_batch_1.json", {
                "_meta": {"asins": ["B0AAA", "B0BBB"]},
                "data": {"items": [{"keyword": "alpha keyword", "searches": 1000, "trafficPercentage": 0.2}]},
            })
            run_script("aggregate_batches.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "keyword_master.json").read_text(encoding="utf-8"))
            rec = result["keywords"][0]
            self.assertEqual(rec["source_asins"], [])
            self.assertEqual(rec["sources"][0]["batch_asins"], ["B0AAA", "B0BBB"])

    def test_keyword_pool_reads_three_pages_deduplicates_all_and_ignores_page4(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            write_json(wd / "input.json", {
                "groups": [{"marketplace": "US", "asins": ["B0AAA"], "node_ids": []}],
                "options": {"data_month": "202607", "keyword_pool_pages": 3},
            })
            page1 = [{"keyword": f"keyword {i:03d}", "searches": 1000 - i} for i in range(50)]
            page2 = ([{"keyword": "keyword 049", "searches": 951}]
                     + [{"keyword": f"keyword {i:03d}", "searches": 1000 - i} for i in range(50, 99)])
            page3 = ([{"keyword": "keyword 098", "searches": 902}]
                     + [{"keyword": f"keyword {i:03d}", "searches": 1000 - i} for i in range(99, 148)])
            write_json(wd / "raw" / "traffic_batch_1_p1.json", {"data": {"items": page1}})
            write_json(wd / "raw" / "traffic_batch_1_p2.json", {"data": {"items": page2}})
            write_json(wd / "raw" / "traffic_batch_1_p3.json", {"data": {"items": page3}})
            write_json(wd / "raw" / "traffic_batch_1_p4.json", {"data": {"items": [
                {"keyword": "must be ignored", "searches": 999999}]}})
            run_script("aggregate_batches.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "keyword_master.json").read_text(encoding="utf-8"))
            self.assertEqual(result["keyword_pool_pages"], 3)
            self.assertEqual(result["keyword_pool_raw_rows"], 150)
            self.assertEqual(result["keyword_pool_cross_page_duplicates"], 2)
            self.assertEqual(result["keyword_pool_actual"], 148)
            self.assertEqual(len(result["keywords"]), 148)
            self.assertNotIn("must be ignored", {r["keyword"] for r in result["keywords"]})

    def test_cross_category_classifier_blocks_modified_core_heads(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            rows = [
                {"keyword": "lunch box", "word_count": 2, "demand_tags": ["core_category:lunch_box"],
                 "relevance": "A", "relevance_reason": "品类词", "purchase_intent": "high"},
                {"keyword": "bento box", "word_count": 2, "demand_tags": ["product_form:bento"],
                 "relevance": "A", "relevance_reason": "产品形态", "purchase_intent": "high"},
                {"keyword": "lunch box boys", "word_count": 3,
                 "demand_tags": ["core_category:lunch_box", "audience:boys"],
                 "relevance": "A", "relevance_reason": "人群", "purchase_intent": "high"},
                {"keyword": "insulated lunch box", "word_count": 3,
                 "demand_tags": ["core_category:lunch_box", "function:insulated"],
                 "relevance": "A", "relevance_reason": "功能", "purchase_intent": "high"},
                {"keyword": "school lunch box", "word_count": 3,
                 "demand_tags": ["core_category:lunch_box", "scene:school"],
                 "relevance": "A", "relevance_reason": "场景", "purchase_intent": "high"},
                {"keyword": "rattan proofing basket", "word_count": 3,
                 "demand_tags": ["core_category:proofing_basket", "material:rattan"],
                 "relevance": "A", "relevance_reason": "材料", "purchase_intent": "high"},
                {"keyword": "girls summer dress", "word_count": 3,
                 "demand_tags": ["core_category:dress", "audience:girls", "scene:summer"],
                 "relevance": "A", "relevance_reason": "人群场景", "purchase_intent": "high"},
                {"keyword": "nike lunch box girls", "word_count": 4, "is_brand": True,
                 "demand_tags": ["brand:nike", "audience:girls"],
                 "relevance": "A", "relevance_reason": "品牌人群", "purchase_intent": "high"},
                {"keyword": "lunch stuff for kids", "word_count": 4,
                 "demand_tags": ["broad:generic", "audience:kids"],
                 "relevance": "B", "relevance_reason": "泛相关人群", "purchase_intent": "medium"},
            ]
            write_json(wd / "cache" / "keyword_labeled.json", {"keywords": rows})
            run_script("classify_keywords.py", "--workdir", wd)
            run_script("validate_labels.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "keyword_labeled.json").read_text(encoding="utf-8"))
            actual = {r["keyword"]: r["traffic_type"] for r in result["keywords"]}
            self.assertEqual(actual["lunch box"], "core_head")
            self.assertEqual(actual["bento box"], "precise_category")
            self.assertEqual(actual["lunch box boys"], "audience_longtail")
            self.assertEqual(actual["insulated lunch box"], "attribute_longtail")
            self.assertEqual(actual["school lunch box"], "scene_longtail")
            self.assertEqual(actual["rattan proofing basket"], "attribute_longtail")
            self.assertEqual(actual["girls summer dress"], "audience_longtail")
            self.assertEqual(actual["nike lunch box girls"], "brand_word")
            self.assertEqual(actual["lunch stuff for kids"], "broad_traffic")

    def test_longtail_can_share_base_family_with_selected_core(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            rows = [
                {"keyword": "lunch box", "normalized_key": "lunch box", "family": "lunch_box",
                 "word_count": 2, "demand_tags": ["core_category:lunch_box"], "traffic_type": "core_head",
                 "relevance": "A", "purchase_intent": "高", "searches": 100000,
                 "purchaseRate": 0.02, "is_brand": False, "is_noise": False},
                {"keyword": "small lunch box", "normalized_key": "small lunch box", "family": "lunch_box",
                 "word_count": 3, "demand_tags": ["core_category:lunch_box", "spec:small"],
                 "traffic_type": "attribute_longtail", "relevance": "A", "purchase_intent": "高",
                 "searches": 20000, "purchaseRate": 0.10, "is_brand": False, "is_noise": False},
            ]
            write_json(wd / "cache" / "keyword_labeled.json", {"keywords": rows})
            run_script("select_trend_keywords.py", "--workdir", wd, "--core", 1, "--longtail", 1, "--cap", 2)
            result = json.loads((wd / "cache" / "trend_keywords.json").read_text(encoding="utf-8"))
            self.assertEqual([r["keyword"] for r in result["keywords"]], ["lunch box", "small lunch box"])

    def test_continuous_trend_window_preserves_internal_gap(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            months = month_range(25, (2024, 7))
            write_json(wd / "cache" / "trend_keywords.json", {
                "keywords": [{"keyword": "alpha", "type": "core", "reason": "fixture"},
                             {"keyword": "beta", "type": "longtail", "reason": "fixture"}]})
            alpha = [{"label": m, "searches": 100 + i, "rank": 1000 - i}
                     for i, m in enumerate(months) if m != "2025-03"]
            beta = [{"label": m, "searches": 200 + i, "rank": 2000 - i}
                    for i, m in enumerate(months)]
            write_json(wd / "raw" / "trend_alpha.json", {"keyword": "alpha", "data": alpha})
            write_json(wd / "raw" / "trend_beta.json", {"keyword": "beta", "data": beta})
            run_script("analyze_trends.py", "--workdir", wd)
            result = json.loads((wd / "cache" / "trend_analysis.json").read_text(encoding="utf-8"))
            self.assertEqual(len(result["months"]), 24)
            self.assertIn("2025-03", result["months"])
            alpha_by_month = {p["month"]: p for p in result["keywords"]["alpha"]["series"]}
            self.assertIsNone(alpha_by_month["2025-03"]["searches"])
            self.assertIn("2025-03", result["keywords"]["alpha"]["stats"]["missing_months"])

    def test_report_two_matrices_and_chart_series(self):
        with tempfile.TemporaryDirectory() as tmp:
            wd = Path(tmp)
            months = month_range()
            keywords = ["alpha keyword", "beta keyword"]
            write_json(wd / "input.json", {
                "groups": [{"marketplace": "US", "asins": ["B0AAA", "B0BBB"], "node_ids": []}],
                "options": {"data_month": "202607", "trend_months": 24, "keyword_pool_pages": 3},
            })
            write_json(wd / "plan.json", {"planned_total": 7, "hard_cap": 9})
            labeled_rows = []
            for i, kw in enumerate(keywords):
                labeled_rows.append({
                    "keyword": kw, "normalized_key": kw, "keywordCn": f"关键词{i + 1}",
                    "translation_fixed": f"关键词{i + 1}", "family": kw,
                    "demand_tags": ["core_category:test"], "traffic_type": "core_head",
                    "traffic_type_reason": "无属性修饰的短品类头部词",
                    "relevance": "A", "relevance_reason": "fixture", "purchase_intent": "高",
                    "searches": 1000 - i * 100, "purchases": 100 - i * 10, "purchaseRate": 0.1,
                    "monopolyClickRate": 0.2 + i * 0.01, "top3ConversionRate": 0.15 + i * 0.01,
                    "bid": 1.2, "bidMin": 1.0, "bidMax": 1.4, "products": 100,
                    "latest1daysAds": 10, "latest30daysAds": 50, "supplyDemandRatio": 2.0,
                    "titleDensity": 1, "source_asins": [], "quality_notes": [],
                })
            write_json(wd / "cache" / "keyword_labeled.json", {"keyword_count": 2, "keywords": labeled_rows})
            trend_keywords = {}
            for i, kw in enumerate(keywords):
                series = [{"month": m, "searches": 1000 + j * 10 + i * 100,
                           "aba_rank": None if (i == 0 and j == 5) else 100 - j - i * 10,
                           "mom": None, "yoy": None,
                           "seasonality_index": 1.0} for j, m in enumerate(months)]
                trend_keywords[kw] = {
                    "type": "core" if i == 0 else "longtail", "series": series,
                    "stats": {"avg_24m": 1100, "median_24m": 1100, "latest_searches": 1200,
                              "mom": 0.01, "yoy": 0.1, "growth_3m": 0.03, "growth_6m": 0.06,
                              "yoy_cumulative_12m": 0.1, "cv": 0.1, "peak_trough_ratio": 1.2,
                              "peak_month": months[-1], "trough_month": months[0],
                              "consecutive_growth_months": 3, "first_growth_inflection": months[5],
                              "two_year_repeat_score": 0.8, "missing_months": []},
                }
            write_json(wd / "cache" / "trend_analysis.json", {
                "months": months, "end_month": months[-1], "window": 24, "keywords": trend_keywords})
            write_json(wd / "cache" / "analysis.json", {
                "category_label": "测试品类", "lifecycle": {"classification": "growth"},
                "entry_conclusion": {"grade": "GO"},
                "risk_checklist": {"cannot_prove": [], "next_validations": []},
                "seasonality": {"timeline_advice": "按相对提前量准备"},
                "demand_map": [{"tag": "core_category:test", "kw_count": 2, "representative": keywords,
                                "searches_median": 950, "purchase_rate_median": 0.1,
                                "click_conc_median": 0.205, "conv_conc_median": 0.155,
                                "ppc_median": 1.2, "trend_direction": "上升", "product_implication": "fixture"}],
                "quadrant": {"click_median": 0.205, "conv_median": 0.155,
                             "items": [{"keyword": keywords[0], "click": 0.2, "conv": 0.15, "interpretation": "fixture"},
                                       {"keyword": keywords[1], "click": 0.21, "conv": 0.16, "interpretation": "fixture"}]},
                "opportunity_matrix": [{"keyword": keywords[0], "type": "core", "evidence": {"searches": 1000}, "confidence": "high"},
                                       {"keyword": keywords[1], "type": "longtail", "evidence": {"searches": 900}, "confidence": "high"}],
                "product_hypotheses": [{
                    "demand": "core_category:test", "source_keywords": [keywords[0]],
                    "evidence": {"searches": 1000, "purchase_rate": 0.1},
                    "user_problem": "用户可能需要明确品类形态（推断）", "product_form": "测试品类",
                    "material": "", "function": "", "spec": "", "bundle_accessory": "",
                    "product_response": "验证测试品类基础形态（假设）", "confidence": "high",
                    "inference_level": "hypothesis", "validation_needed": "评论与供应链验证",
                }],
            })
            calls = [
                {"seq": 1, "tool": "traffic_extend", "scope": "US batch=1 page=1 historyDate=202607", "success": True},
                {"seq": 2, "tool": "traffic_extend", "scope": "US batch=1 page=2 historyDate=202607", "success": True},
                {"seq": 3, "tool": "traffic_extend", "scope": "US batch=1 page=3 historyDate=202607", "success": True},
                {"seq": 4, "tool": "traffic_keyword", "scope": "B0AAA page=1 current", "success": True},
                {"seq": 5, "tool": "traffic_keyword", "scope": "B0BBB page=1 current", "success": True},
                {"seq": 6, "tool": "aba_research_trend", "scope": keywords[0], "success": True},
                {"seq": 7, "tool": "aba_research_trend", "scope": keywords[1], "success": True},
            ]
            write_json(wd / "cache" / "call_log.json", {"planned_total": 7, "hard_cap": 9, "actual_total": 7, "calls": calls})
            write_json(wd / "cache" / "asin_keyword_matrix.json", {
                "schema_version": 2, "marketplace": "US", "snapshot_type": "current",
                "scope": "current_top100", "page_limit": 1, "page_size": 100,
                "observed_at": "2026-08-27T10:00:00+08:00", "asins": ["B0AAA", "B0BBB"],
                "coverage": {"B0AAA": {"request_success": True}, "B0BBB": {"request_success": True}},
                "records": [
                    {"asin": "B0AAA", "keyword": keywords[0], "normalized_key": keywords[0],
                     "natural_page": 1, "natural_index": 36, "natural_rank": 36,
                     "ad_page": 2, "ad_index": 8, "ad_rank": 96, "traffic_percentage": 0.13},
                    {"asin": "B0BBB", "keyword": keywords[0], "normalized_key": keywords[0],
                     "natural_page": 1, "natural_index": 31, "natural_rank": 31,
                     "ad_page": None, "ad_index": None, "ad_rank": None, "traffic_percentage": 0.03},
                    {"asin": "B0AAA", "keyword": keywords[1], "normalized_key": keywords[1],
                     "natural_page": None, "natural_index": None, "natural_rank": None,
                     "ad_page": 1, "ad_index": 4, "ad_rank": 4, "traffic_percentage": None},
                ],
            })

            run_script("build_excel_report.py", "--workdir", wd)
            report = next((wd / "output").glob("*.xlsx"))
            wb = load_workbook(report)
            self.assertEqual(wb.sheetnames, [
                "00_市场看板", "01_关键词总表", "02_需求分类汇总", "03_24月搜索量_宽表",
                "04_24月搜索量_长表", "05_趋势与季节性", "06_集中度与机会矩阵",
                "07_产品方向", "08_ASIN自然广告位矩阵", "09_ASIN流量占比矩阵",
                "10_市场进入与风险", "11_调研范围与调用", "12_数据质量"])
            rank_ws = wb["08_ASIN自然广告位矩阵"]
            traffic_ws = wb["09_ASIN流量占比矩阵"]
            self.assertEqual([rank_ws.cell(1, c).value for c in range(1, 6)], [
                "矩阵观测时间", "原始关键词", "中文翻译",
                "B0AAA|自然位/广告位", "B0BBB|自然位/广告位"])
            self.assertEqual(rank_ws["D2"].value, "P1-36/P2-8")
            self.assertEqual(rank_ws["E2"].value, "P1-31/")
            self.assertEqual(rank_ws["D3"].value, "/P1-4")
            self.assertAlmostEqual(traffic_ws["D2"].value, 0.13)
            self.assertIsNone(traffic_ws["D3"].value)

            dash = wb["00_市场看板"]
            self.assertEqual(sorted(type(c).__name__ for c in dash._charts),
                             ["BarChart", "LineChart", "LineChart", "ScatterChart"])
            line_charts = [c for c in dash._charts if isinstance(c, LineChart)]
            self.assertEqual([[series_title(s) for s in c.series] for c in line_charts],
                             [keywords, keywords])
            self.assertEqual([series_color(s) for s in line_charts[0].series],
                             [series_color(s) for s in line_charts[1].series])
            self.assertEqual(
                [series_category_formula(s) for s in line_charts[0].series],
                ["'03_24月搜索量_宽表'!$D$1:$AA$1"] * 2,
            )
            self.assertEqual(
                [series_category_formula(s) for s in line_charts[1].series],
                ["'04_24月搜索量_长表'!$E$2:$E$25",
                 "'04_24月搜索量_长表'!$E$26:$E$49"],
            )
            self.assertEqual([c.x_axis.tickLblSkip for c in line_charts], [2, 2])
            self.assertEqual(line_charts[1].y_axis.scaling.orientation, "minMax")
            scatter = next(c for c in dash._charts if type(c).__name__ == "ScatterChart")
            self.assertFalse(scatter.visible_cells_only)
            self.assertEqual(scatter.scatterStyle, "marker")
            self.assertEqual((scatter.x_axis.numFmt.formatCode, scatter.y_axis.numFmt.formatCode), ("0%", "0%"))
            self.assertEqual([series_title(s) for s in scatter.series[:1]], ["核心/精准词"])
            self.assertNotEqual((scatter.x_axis.scaling.min, scatter.x_axis.scaling.max), (0, 1))
            self.assertGreaterEqual(scatter.x_axis.scaling.max - scatter.x_axis.scaling.min, 0.15)
            self.assertIn("点击中位线", [series_title(s) for s in scatter.series])
            self.assertIn("转化中位线", [series_title(s) for s in scatter.series])
            run_script("validate_report.py", "--workdir", wd, "--file", report)


if __name__ == "__main__":
    unittest.main()
